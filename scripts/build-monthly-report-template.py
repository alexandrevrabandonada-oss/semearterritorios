"""Generate the editable monthly template using the same headings as the form."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
source = (root / "lib/professional-monthly-reports.ts").read_text(encoding="utf-8")
fields = re.findall(r'\["([a-z_]+)", "([^"]+)"\]', source.split("] as const;")[0])
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.85)
for name in ("Normal", "Title", "Heading 1"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
doc.styles["Title"].font.size = Pt(23)
doc.styles["Heading 1"].font.size = Pt(13)
doc.add_paragraph("Relatório mensal do profissional", "Title")
doc.add_paragraph("SEMEAR Territórios • Volta Redonda RJ")
doc.add_paragraph("Este relatório apresenta as atividades e entregas do mês para revisão técnica da coordenação do projeto de acompanhamento acadêmico e popular da poluição do ar. Preencha os campos e mantenha os títulos para facilitar a importação no sistema.")
doc.add_paragraph("Profissional responsável [preencher]\nMês e ano de referência [preencher]\nData da entrega [preencher]")
prompts = {
    "summary": "Resuma o trabalho realizado e os principais resultados do mês.",
    "objectives": "Informe os objetivos de educação ambiental popular, cartografia social e diagnóstico participativo trabalhados no período.",
    "activities": "Liste datas, ações, oficinas, territórios e participação aproximada. Indique as entregas e os materiais produzidos. Caso não tenha havido atividade de campo, justifique.",
    "methodologies": "Descreva dinâmicas, atividades formativas, cartografia, diagnóstico e plantio utilizados.",
    "participant_accounts": "Apresente percepções coletivas, consensos e divergências, sem identificar moradores.",
    "team_observations": "Descreva o que a equipe observou, distinguindo observações de relatos e medições.",
    "research_references": "Identifique pesquisas, medições e fontes consultadas. Se não houve medição, informe. Não atribua causalidade de doenças a partir de percepções.",
    "learnings": "Registre aprendizados para o projeto e para as próximas oficinas.",
    "difficulties": "Registre desafios operacionais, limites metodológicos e lacunas do diagnóstico.",
    "proposals": "Separe propostas apresentadas de acordos pactuados e ações concluídas. Indique responsáveis coletivos quando definidos.",
    "pending_items": "Informe o que depende de revisão, decisão ou continuidade.",
    "next_steps": "Descreva as atividades previstas e os próximos passos, sem apresentar propostas não pactuadas como compromissos.",
}
for index, (key, label) in enumerate(fields):
    if index in (4, 9):
        doc.add_page_break()
    doc.add_paragraph(label, "Heading 1")
    doc.add_paragraph(prompts[key])
    doc.add_paragraph("[Preencher]")
doc.add_paragraph("Materiais e cuidados de envio", "Heading 1")
doc.add_paragraph("No sistema, classifique as entregas e vincule seus anexos às atividades. Para mapas sociais, informe o contexto de construção e a legenda. Envie este documento em DOCX ou PDF digital e selecione-o como principal. Não inclua CPF, telefone pessoal, endereço residencial ou dados identificáveis de saúde de moradores.")
doc.core_properties.title = "Relatório mensal do profissional"
doc.core_properties.author = "SEMEAR Territórios"
# The bundled Word template can carry a blue border in the Title style.
for element in (doc.element, doc.styles.element):
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
output = root / "public/modelos/relatorio-mensal-profissional.docx"
output.parent.mkdir(parents=True, exist_ok=True)
doc.save(output)
print(output)
